"""
utils/text_extractor.py

Handles extracting raw text from user-supplied educational documents.
Supported formats: PDF, DOCX, TXT.

Each extractor is defensive: it never raises a raw exception up to the
Streamlit UI. Instead it raises a `TextExtractionError` with a clean,
user-friendly message, which app.py catches and displays.
"""

from __future__ import annotations

import re


class TextExtractionError(Exception):
    """Raised when text cannot be extracted from a document."""
    pass


def _clean_text(text: str) -> str:
    """Normalize whitespace so downstream NLP has consistent input."""
    if not text:
        return ""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text_from_pdf(file) -> str:
    """Extract text from a PDF file (file-like object)."""
    try:
        from pypdf import PdfReader
    except ImportError:
        try:
            from PyPDF2 import PdfReader  # fallback for older envs
        except ImportError as exc:
            raise TextExtractionError(
                "PDF support is not installed. Please install 'pypdf'."
            ) from exc

    try:
        file.seek(0)
    except Exception:
        pass

    try:
        reader = PdfReader(file)
    except Exception as exc:
        raise TextExtractionError(
            "This PDF could not be opened. It may be corrupted or password protected."
        ) from exc

    pages_text = []
    for page in reader.pages:
        try:
            page_text = page.extract_text() or ""
        except Exception:
            page_text = ""
        pages_text.append(page_text)

    full_text = _clean_text("\n".join(pages_text))

    if not full_text:
        raise TextExtractionError(
            "No readable text was found in this PDF. It may be a scanned "
            "image without a text layer."
        )

    return full_text


def extract_text_from_docx(file) -> str:
    """Extract text from a DOCX file (file-like object)."""
    try:
        import docx
    except ImportError as exc:
        raise TextExtractionError(
            "DOCX support is not installed. Please install 'python-docx'."
        ) from exc

    try:
        file.seek(0)
    except Exception:
        pass

    try:
        document = docx.Document(file)
    except Exception as exc:
        raise TextExtractionError(
            "This DOCX file could not be opened. It may be corrupted or "
            "in an unsupported format."
        ) from exc

    paragraphs = [p.text for p in document.paragraphs if p.text and p.text.strip()]

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text and cell.text.strip():
                    paragraphs.append(cell.text.strip())

    full_text = _clean_text("\n".join(paragraphs))

    if not full_text:
        raise TextExtractionError("No readable text was found in this DOCX file.")

    return full_text


def extract_text_from_txt(file) -> str:
    """Extract text from a TXT file (file-like object)."""
    try:
        file.seek(0)
    except Exception:
        pass

    try:
        raw = file.read()
    except Exception as exc:
        raise TextExtractionError("This TXT file could not be read.") from exc

    if isinstance(raw, bytes):
        for encoding in ("utf-8", "utf-8-sig", "latin-1"):
            try:
                raw = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise TextExtractionError("This TXT file's encoding could not be detected.")

    full_text = _clean_text(raw)

    if not full_text:
        raise TextExtractionError("This TXT file appears to be empty.")

    return full_text


def extract_text(uploaded_file) -> str:
    """Dispatch to the correct extractor based on file extension."""
    if uploaded_file is None:
        raise TextExtractionError("No file was provided.")

    name = getattr(uploaded_file, "name", "") or ""
    extension = name.lower().rsplit(".", 1)[-1] if "." in name else ""

    if extension == "pdf":
        return extract_text_from_pdf(uploaded_file)
    elif extension == "docx":
        return extract_text_from_docx(uploaded_file)
    elif extension == "txt":
        return extract_text_from_txt(uploaded_file)
    else:
        raise TextExtractionError(
            f"Unsupported file type: '.{extension}'. Please upload a "
            f"PDF, DOCX, or TXT file."
        )
